from __future__ import annotations

import json
import re
import hashlib
from pathlib import Path
from typing import Any, Iterator

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P


# =========================
# 一、基础配置
# =========================

INPUT_DIR = Path("./docs")
OUTPUT_DIR = Path("./output")
OUTPUT_FILE = OUTPUT_DIR / "chunks.jsonl"

# 普通文本块的目标字符数
CHUNK_SIZE = 1000

# 只有普通文本块允许合并，不切开代码和表格
MIN_CHUNK_SIZE = 250

# 这些内容通常出现在单元格首行，表示这是代码块
CODE_LANGUAGES = {
    "bash", "shell", "sh", "python", "python3",
    "javascript", "java", "c++", "cpp", "c",
    "xml", "yaml", "yml", "json", "cmake",
    "plain text", "plaintext", "text",
    "html", "css", "sql", "dockerfile",
    "powershell", "markdown",
}

# 章节标题，例如：
# 1.2 ROS 核心概念
# 2.3.1 创建话题发布者文件
NUMBERED_HEADING = re.compile(
    r"^\s*(\d+(?:\.\d+)*\.?)\s+(.+?)\s*$"
)


# =========================
# 二、按原始顺序遍历 DOCX
# =========================

def iter_block_items(parent) -> Iterator[Any]:
    """
    按照文档中的原始顺序遍历段落和表格。

    不能只遍历 doc.paragraphs：
    Word 中的代码、命令和知识表格经常存放在表格里。
    """
    if hasattr(parent, "element") and hasattr(
        parent.element, "body"
    ):
        parent_element = parent.element.body
    else:
        parent_element = parent._tc

    for child in parent_element.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)

        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)


# =========================
# 三、文本和表格处理
# =========================

def clean_text(text: str) -> str:
    """清理普通文本中的多余空白，保留必要换行。"""
    text = text.replace("\xa0", " ")
    text = text.replace("\u3000", " ")
    lines = [line.strip() for line in text.splitlines()]
    return "\n".join(
        line for line in lines if line
    ).strip()


def escape_markdown_cell(text: str) -> str:
    """避免表格单元格中的竖线破坏 Markdown 表格结构。"""
    return (
        text.replace("\\", "\\\\")
        .replace("|", "\\|")
        .replace("\n", "<br>")
    )


def table_to_markdown(table: Table) -> str:
    """
    将 Word 表格转换为 Markdown。

    保留行列关系、表头及单元格内容。
    默认把第一行视为表头。
    """
    rows = []

    for row in table.rows:
        cells = [
            escape_markdown_cell(
                clean_text(cell.text)
            )
            for cell in row.cells
        ]

        # Word 中的合并单元格可能重复出现
        # 这里保留实际读取到的列结构，不擅自删除重复值
        rows.append(cells)

    if not rows:
        return ""

    column_count = max(len(row) for row in rows)

    # 统一列数，避免 Markdown 表格格式不完整
    rows = [
        row + [""] * (column_count - len(row))
        for row in rows
    ]

    header = rows[0]
    separator = ["---"] * column_count

    output = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join(separator) + " |",
    ]

    for row in rows[1:]:
        output.append(
            "| " + " | ".join(row) + " |"
        )

    return "\n".join(output)


def detect_code_table(table: Table):
    """
    检测常见的 Word 代码块。

    例如单个单元格的内容：
        Bash
        $ roscore

    返回 (language, code)，否则返回 None。
    """
    if len(table.rows) != 1 or len(table.columns) != 1:
        return None

    text = table.cell(0, 0).text.strip()
    if not text:
        return None

    lines = text.splitlines()
    first_line = lines[0].strip()

    language = first_line.lower()

    if language in CODE_LANGUAGES and len(lines) > 1:
        code = "\n".join(lines[1:]).strip("\n")
        if code:
            return language, code

    # 对没有语言标签的单元格做保守判断。
    # 避免将普通知识段落误识别为代码。
    if len(lines) >= 2:
        code_signals = (
            lines[0].lstrip().startswith((
                "$ ", "#!/", "import ", "from ",
                "def ", "class ", "sudo ",
                "rosrun ", "rosservice ",
                "rostopic ", "rosparam ",
                "catkin_make", "find_package(",
                "<build_depend>", "<exec_depend>",
            ))
        )

        if code_signals:
            return "text", text

    return None


# =========================
# 四、识别标题与知识单元
# =========================

def detect_heading(paragraph: Paragraph):
    """
    识别标题。

    优先使用 Word 的标题样式；
    对全部采用 Normal 样式的教程，再识别编号标题。
    """
    text = clean_text(paragraph.text)
    if not text:
        return None

    style_name = paragraph.style.name.lower()

    if style_name.startswith("heading"):
        match = re.search(r"(\d+)", style_name)
        level = int(match.group(1)) if match else 1
        return level, text

    match = NUMBERED_HEADING.match(text)

    if match:
        number = match.group(1).rstrip(".")
        title = match.group(2).strip()

        # 只把较短、具有编号结构的段落识别为标题
        # 避免误判正文中以数字开头的句子
        if len(text) <= 100:
            level = number.count(".") + 1
            return level, text

    return None


def extract_document(docx_path: Path) -> list[dict]:
    """
    抽取文档中的知识单元。

    每个单元保存：
    - content: 实际内容
    - type: 内容类型
    - headings: 所属标题路径
    - source: 来源文件
    - sequence: 原始顺序
    """
    doc = Document(str(docx_path))

    blocks = []
    heading_stack = []

    for item in iter_block_items(doc):
        if isinstance(item, Paragraph):
            text = clean_text(item.text)
            if not text:
                continue

            heading = detect_heading(item)

            if heading:
                level, heading_text = heading

                # 更新标题层级
                heading_stack = [
                    entry for entry in heading_stack
                    if entry["level"] < level
                ]

                heading_stack.append({
                    "level": level,
                    "title": heading_text,
                })

                blocks.append({
                    "type": "heading",
                    "content": heading_text,
                    "headings": [
                        entry["title"]
                        for entry in heading_stack
                    ],
                    "source": docx_path.name,
                    "sequence": len(blocks),
                })
                continue

            blocks.append({
                "type": "text",
                "content": text,
                "headings": [
                    entry["title"]
                    for entry in heading_stack
                ],
                "source": docx_path.name,
                "sequence": len(blocks),
            })

        elif isinstance(item, Table):
            code_info = detect_code_table(item)

            if code_info:
                language, code = code_info
                block_type = (
                    "command"
                    if language in {"bash", "shell", "sh"}
                    else "code"
                )

                blocks.append({
                    "type": block_type,
                    "language": language,
                    "content": code,
                    "headings": [
                        entry["title"]
                        for entry in heading_stack
                    ],
                    "source": docx_path.name,
                    "sequence": len(blocks),
                })
            else:
                table_text = table_to_markdown(item)

                if table_text.strip():
                    blocks.append({
                        "type": "table",
                        "content": table_text,
                        "headings": [
                            entry["title"]
                            for entry in heading_stack
                        ],
                        "source": docx_path.name,
                        "sequence": len(blocks),
                    })

    return blocks


# =========================
# 五、生成适合 RAG 的分块
# =========================

def make_chunk(
    blocks: list[dict],
    chunk_type: str,
) -> dict:
    """为一个最终分块生成稳定的 ID 和元数据。"""
    first = blocks[0]
    last = blocks[-1]

    headings = first["headings"]

    content_parts = []

    # 将标题路径写入正文，帮助检索定位上下文
    if headings:
        content_parts.append(
            "所属章节：" + " > ".join(headings)
        )

    for block in blocks:
        if block["type"] == "heading":
            continue

        content_parts.append(block["content"])

    content = "\n\n".join(content_parts).strip()

    identity = (
        first["source"]
        + str(first["sequence"])
        + str(last["sequence"])
        + content
    )

    chunk_id = hashlib.sha256(
        identity.encode("utf-8")
    ).hexdigest()[:20]

    return {
        "id": chunk_id,
        "content": content,
        "metadata": {
            "source": first["source"],
            "headings": headings,
            "heading_path": " > ".join(headings),
            "type": chunk_type,
            "start_sequence": first["sequence"],
            "end_sequence": last["sequence"],
        },
    }


def chunk_document(
    blocks: list[dict],
    chunk_size: int = CHUNK_SIZE,
    min_chunk_size: int = MIN_CHUNK_SIZE,
) -> list[dict]:
    """
    结构感知分块：

    1. 表格独立成块。
    2. 命令和代码独立成块。
    3. 普通文本按字符数合并。
    4. 遇到标题时刷新当前文本块。
    5. 不跨越特殊内容合并普通文本。
    """
    chunks = []
    text_buffer = []
    text_length = 0

    def flush_text():
        nonlocal text_buffer, text_length

        if not text_buffer:
            return

        chunks.append(
            make_chunk(text_buffer, "text")
        )

        text_buffer = []
        text_length = 0

    for block in blocks:
        block_type = block["type"]

        if block_type == "heading":
            flush_text()
            continue

        content = block["content"]

        # 特殊内容作为独立知识单元
        if block_type in {
            "table", "code", "command"
        }:
            flush_text()
            chunks.append(
                make_chunk([block], block_type)
            )
            continue

        # 普通文本只在字符数范围内合并
        if (
            text_buffer
            and text_length + len(content) > chunk_size
        ):
            flush_text()

        text_buffer.append(block)
        text_length += len(content)

    flush_text()

    # 尝试合并过短的普通文本块。
    # 不合并表格、代码或命令。
    merged = []

    for chunk in chunks:
        if (
            merged
            and chunk["metadata"]["type"] == "text"
            and merged[-1]["metadata"]["type"] == "text"
            and len(merged[-1]["content"]) < min_chunk_size
            and (
                len(merged[-1]["content"])
                + len(chunk["content"])
                <= chunk_size
            )
            and merged[-1]["metadata"]["heading_path"]
            == chunk["metadata"]["heading_path"]
        ):
            previous = merged.pop()

            combined_content = (
                previous["content"]
                + "\n\n"
                + chunk["content"]
            )

            combined_id = hashlib.sha256(
                (
                    previous["id"]
                    + chunk["id"]
                    + combined_content
                ).encode("utf-8")
            ).hexdigest()[:20]

            previous["content"] = combined_content
            previous["id"] = combined_id
            previous["metadata"]["end_sequence"] = (
                chunk["metadata"]["end_sequence"]
            )

            merged.append(previous)
        else:
            merged.append(chunk)

    return merged


# =========================
# 六、批量处理与 JSONL 输出
# =========================

def process_docx(
    docx_path: Path,
    chunk_size: int = CHUNK_SIZE,
) -> list[dict]:
    blocks = extract_document(docx_path)

    chunks = chunk_document(
        blocks,
        chunk_size=chunk_size,
    )

    return chunks


def main():
    OUTPUT_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    files = sorted(INPUT_DIR.glob("*.docx"))

    if not files:
        print(f"未找到 DOCX 文件：{INPUT_DIR.resolve()}")
        return

    total_chunks = 0

    with OUTPUT_FILE.open(
        "w",
        encoding="utf-8",
    ) as output:
        for docx_path in files:
            try:
                chunks = process_docx(docx_path)

                for chunk in chunks:
                    output.write(
                        json.dumps(
                            chunk,
                            ensure_ascii=False,
                        )
                        + "\n"
                    )

                total_chunks += len(chunks)

                print(
                    f"{docx_path.name}: "
                    f"{len(chunks)} 个分块"
                )

            except Exception as exc:
                print(
                    f"处理失败：{docx_path.name}: {exc}"
                )

    print(f"总分块数：{total_chunks}")
    print(f"输出文件：{OUTPUT_FILE.resolve()}")


if __name__ == "__main__":
    main()
